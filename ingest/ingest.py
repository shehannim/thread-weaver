import os
import json
import pytesseract
from PIL import Image
import fitz  # PyMuPDF
import pdfplumber
import docx
from langchain_text_splitters import RecursiveCharacterTextSplitter
from dotenv import load_dotenv

load_dotenv()

TESSERACT_CMD = os.getenv("TESSERACT_CMD")
if TESSERACT_CMD:
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD

def get_doc_metadata(filepath):
    path_lower = filepath.lower()
    if "codex" in path_lower:
        return "codex", "high"
    elif "wiki" in path_lower:
        return "wiki", "medium"
    elif "novel" in path_lower:
        return "novel", "medium"
    elif "ephemera" in path_lower:
        return "ephemera", "low"
    else:
        return "unknown", "low"

def chunk_text(
    text,
    doc_id,
    doc_type,
    reliability,
    start_page,
    source_path=None,
    extraction_method="native_text"
):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=150,
        separators=["\n\n", "\n", ".", " ", ""]
    )
    chunks = splitter.split_text(text)
    
    records = []
    for i, chunk in enumerate(chunks):
        cleaned_text = chunk.strip()
        if not cleaned_text:
            continue
        records.append({
            "doc_id": doc_id,
            "doc_type": doc_type,
            "source_reliability": reliability,
            "page_number": start_page,
            "chunk_text": cleaned_text,
            "chunk_id": f"{doc_id}__p{start_page}__c{i}",
            "source_path": source_path,
            "extraction_method": extraction_method,
            "chunk_index": i
        })
    return records

def process_pdf(filepath, doc_type, doc_id, reliability, source_path):
    records = []
    if doc_type == "codex":
        with pdfplumber.open(filepath) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                method = "native_text"
                if len(text.strip()) < 50:
                    im = page.to_image(resolution=300)
                    text = pytesseract.image_to_string(im.original)
                    method = "ocr"
                
                if text.strip():
                    records.extend(chunk_text(text, doc_id, doc_type, reliability, i + 1, source_path, method))
    else:
        doc = fitz.open(filepath)
        for i in range(len(doc)):
            page = doc[i]
            text = page.get_text()
            method = "native_text"
            
            if len(text.strip()) < 50:
                pix = page.get_pixmap(dpi=300)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                text = pytesseract.image_to_string(img)
                method = "ocr"
            
            if text.strip():
                records.extend(chunk_text(text, doc_id, doc_type, reliability, i + 1, source_path, method))
        doc.close()
    return records

def process_docx(filepath, doc_id, doc_type, reliability, source_path):
    doc = docx.Document(filepath)
    paras = [para.text.strip() for para in doc.paragraphs if para.text.strip()]

    tables = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                tables.append(" | ".join(cells))

    text = "\n".join(paras + tables)
    if text.strip():
        return chunk_text(text, doc_id, doc_type, reliability, 1, source_path, "native_text")
    return []

def process_text(filepath, doc_id, doc_type, reliability, source_path):
    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        text = f.read()
    if text.strip():
        return chunk_text(text, doc_id, doc_type, reliability, 1, source_path, "native_text")
    return []

def process_image(filepath, doc_id, doc_type, reliability, source_path):
    img = Image.open(filepath)
    text = pytesseract.image_to_string(img)
    if text.strip():
        return chunk_text(text, doc_id, doc_type, reliability, 1, source_path, "ocr")
    return []

def run_ingestion(raw_dir, output_file):
    if not os.path.exists(raw_dir):
        print(f"Directory {raw_dir} does not exist.")
        return

    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    with open(output_file, 'w', encoding='utf-8') as f_out:
        for root, dirs, files in os.walk(raw_dir):
            dirs.sort()
            files.sort()
            for file in files:
                filepath = os.path.join(root, file)

                # Calculate relative path and normalize slashes
                rel_path = os.path.relpath(filepath, raw_dir).replace("\\", "/")

                # Derive doc_id from relative path without extension, replacing slashes with __
                doc_id = os.path.splitext(rel_path)[0].replace("/", "__")

                doc_type, reliability = get_doc_metadata(filepath)
                
                ext = os.path.splitext(file)[1].lower()
                records = []
                try:
                    if ext == '.pdf':
                        records = process_pdf(filepath, doc_type, doc_id, reliability, rel_path)
                    elif ext == '.docx':
                        records = process_docx(filepath, doc_id, doc_type, reliability, rel_path)
                    elif ext in ['.md', '.txt']:
                        records = process_text(filepath, doc_id, doc_type, reliability, rel_path)
                    elif ext in ['.jpg', '.jpeg', '.png']:
                        records = process_image(filepath, doc_id, doc_type, reliability, rel_path)
                    else:
                        print(f"Skipping unsupported file: {filepath}")
                        continue
                except Exception as e:
                    print(f"Error processing {filepath}: {e}")
                
                for record in records:
                    f_out.write(json.dumps(record) + '\n')
                
                print(f"Processed {file} ({len(records)} chunks)")

if __name__ == "__main__":
    CORPUS_PATH = os.getenv("CORPUS_PATH", os.path.join(os.path.dirname(__file__), "../data/raw"))
    OUTPUT_FILE = os.path.join(os.path.dirname(__file__), "../data/processed/chunks.jsonl")
    
    print(f"Starting ingestion from {CORPUS_PATH}...")
    run_ingestion(CORPUS_PATH, OUTPUT_FILE)
    print(f"Ingestion complete. Output saved to {OUTPUT_FILE}")
